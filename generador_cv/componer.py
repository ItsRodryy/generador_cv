"""Native editable Word with one configurable Minimalista Onda template."""
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from .modelo import ROOT, load
from .recursos import wavy_photo, icon, corner

DEFAULT_STYLE='minimalista_onda'

def style(name=DEFAULT_STYLE, overrides=None):
    if name != DEFAULT_STYLE:raise ValueError('La única plantilla disponible es minimalista_onda')
    result=load(ROOT / 'estilos'/f'{name}.json')
    result.update(overrides or {})
    for k in ('margin','body','name','role','photo_pt','header_radius','header_pad','section_gap','block_gap','icon_pt'):
        if k in result and not 0 < float(result[k]) < 100:raise ValueError('Tamaño de estilo inválido: '+k)
    for k in ('header_bg','wave_color','icon_color','divider_color'):
        if len(result[k])!=6 or not all(ch in '0123456789abcdefABCDEF' for ch in result[k]):raise ValueError('Color inválido: '+k)
    return result

def _rgb(hexstr):return RGBColor.from_string(hexstr.upper())

def _para(parent, text='', size=None, bold=False, gray=None, before=0, after=0):
    if hasattr(parent,'_tc') and len(parent.paragraphs)==1 and not parent.paragraphs[0].text and not parent.paragraphs[0].runs:
        p=parent.paragraphs[0]
    else:p=parent.add_paragraph()
    f=p.paragraph_format
    f.space_before=Pt(before);f.space_after=Pt(after);f.widow_control=True;f.keep_together=True
    if text:
        r=p.add_run(str(text));r.bold=bold
        if size:r.font.size=Pt(size)
        if gray:r.font.color.rgb=_rgb(gray)
    return p

def _link(p,label,url,gray='555555'):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');clr=OxmlElement('w:color');clr.set(qn('w:val'),gray);pr.append(clr);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def _shading(cell,color):
    tc=cell._tc.get_or_add_tcPr();sh=tc.find(qn('w:shd'))
    if sh is None:sh=OxmlElement('w:shd');tc.append(sh)
    sh.set(qn('w:fill'),color);sh.set(qn('w:val'),'clear')

def _cell_margin(cell,amount=0):
    tc=cell._tc.get_or_add_tcPr();mar=tc.find(qn('w:tcMar'))
    if mar is None:mar=OxmlElement('w:tcMar');tc.append(mar)
    for side in ('top','start','bottom','end'):
        el=mar.find(qn('w:'+side))
        if el is None:el=OxmlElement('w:'+side);mar.append(el)
        el.set(qn('w:w'),str(int(amount*20)));el.set(qn('w:type'),'dxa')

def _table(parent, rows, cols, widths):
    table=parent.add_table(rows=rows,cols=cols)
    table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,w in enumerate(widths):table.columns[j].width=Pt(w)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'nil');borders.append(el)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        for j,cell in enumerate(row.cells):cell.width=Pt(widths[j]);_cell_margin(cell)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return table

def _header(doc,data,config,look,case,tmp):
    person=data['person'];r=float(look['header_radius'])
    pagewidth=210*72/25.4-2*look['margin']*72/25.4
    outer=_table(doc,3,3,[r,pagewidth-2*r,r])
    for row in outer.rows:
        for c in row.cells:_shading(c,look['header_bg'])
    for ri in (0,2):
        outer.rows[ri].height=Pt(r);outer.rows[ri].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
    for i,j,code in [(0,0,'tl'),(0,2,'tr'),(2,0,'bl'),(2,2,'br')]:
        c=outer.cell(i,j);_shading(c,'FFFFFF');p=c.paragraphs[0]
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=Pt(r)
        p.add_run().add_picture(str(corner(code,tmp/f'c-{code}.png',look)),width=Pt(r),height=Pt(r))
    mid=outer.cell(1,1)
    pad=float(look['header_pad'])
    width=pagewidth-2*r
    # Text and photograph remain independent editable Word content.
    photo_enabled=bool(config.get('foto'))
    available=width-2*pad
    img_w=float(look['photo_pt']) if photo_enabled else 0
    cols=[img_w+9,available-img_w-9] if photo_enabled else [available]
    inner=_table(mid,1,len(cols),cols)
    if photo_enabled:
        photo=case/config['foto']
        if not photo.is_file():raise ValueError('Fotografía no encontrada: '+str(photo))
        p=inner.cell(0,0).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(wavy_photo(photo,tmp/'foto_onda.png',look)),width=Pt(img_w),height=Pt(img_w))
        txt=inner.cell(0,1)
    else:txt=inner.cell(0,0)
    txt.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _para(txt,person['name'],size=look['name'],bold=True,after=2)
    if person.get('headline'):_para(txt,person['headline'],size=look['role'],bold=True,gray='555555',after=3)
    # Compact contacts, preserving URLs and email hyperlinks.
    contacts=person.get('contacts',[])
    for start in range(0,len(contacts),2):
        p=_para(txt,after=1)
        for j,item in enumerate(contacts[start:start+2]):
            if j:p.add_run('  ·  ')
            if item.startswith(('https://','http://','mailto:')):_link(p,item,item)
            elif '@' in item and ' ' not in item:_link(p,item,'mailto:'+item)
            else:p.add_run(item)
        for rr in p.runs:rr.font.size=Pt(look['contact'])
    # Preserve editable table in middle cell.
    # The header is already followed by the first section; no spacer paragraph.

def _section(parent, section, category, look, tmp):
    p=_para(parent,before=look['section_gap'],after=3)
    p.paragraph_format.keep_with_next=True
    if look.get('icons',True):
        f=icon(category,tmp/f'icon_{category}.png',look['icon_color'])
        p.add_run().add_picture(str(f),width=Pt(look['icon_pt']),height=Pt(look['icon_pt']))
        p.add_run('  ')
    rr=p.add_run(section.upper());rr.bold=True;rr.font.size=Pt(look['section']);rr.font.color.rgb=_rgb('333333')
    # bottom rule
    if look.get('separator')=='horizontal':
        pPr=p._p.get_or_add_pPr();bdr=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'5');edge.set(qn('w:color'),look['divider_color']);edge.set(qn('w:space'),'3');bdr.append(edge);pPr.append(bdr)

def _block(doc,b,look,tmp,section_on=False):
    if section_on:_section(doc,b['section'],b.get('category','other'),look,tmp)
    head=_para(doc,b.get('title',''),size=look['title'],bold=True,after=1)
    head.paragraph_format.keep_with_next=True
    if b.get('meta'):
        meta=_para(doc,b['meta'],size=max(9.0,look['body']),gray='666666',after=2)
        meta.paragraph_format.keep_with_next=True
    items=[('body',v) for v in b.get('body',[])]+[('bullet',v) for v in b.get('bullets',[])]
    for ix,(typ,v) in enumerate(items):
        p=_para(doc,('•  ' if typ=='bullet' else '')+v,after=1)
        if typ=='bullet':p.paragraph_format.left_indent=Mm(3);p.paragraph_format.first_line_indent=Mm(-3)
        if ix<len(items)-1:p.paragraph_format.keep_with_next=True
    for item in b.get('links',[]):
        p=_para(doc,after=1);_link(p,item['label'],item['url'])
    # Last real paragraph supplies the block spacing; no blank spacer lines.
    if doc.paragraphs:doc.paragraphs[-1].paragraph_format.space_after=Pt(look['block_gap'])

def compose(data,config,name,path,cuts,case):
    look=style(name,config.get('visual'))
    if cuts is None:groups=[data['blocks']]
    elif isinstance(cuts,int):groups=[data['blocks'][:cuts],data['blocks'][cuts:]]
    else:
        boundaries=[0,*cuts,len(data['blocks'])];groups=[data['blocks'][a:b] for a,b in zip(boundaries,boundaries[1:])]
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.left_margin=sec.right_margin=Mm(look['margin']);sec.top_margin=Mm(look['margin']);sec.bottom_margin=Mm(look['margin'])
    sec.header_distance=Mm(7);sec.footer_distance=Mm(7)
    base=doc.styles['Normal'];base.font.name=config.get('fuente') or 'Lato';base.font.size=Pt(look['body']);base.font.color.rgb=_rgb('333333')
    base.paragraph_format.line_spacing=look['leading']
    blocks=data['blocks'];first_section={};last=None
    for b in blocks:
        if b.get('section'):last=(b['section'],b.get('category','other'))
        first_section[b['id']]=last
    with TemporaryDirectory(prefix='cv-assets-') as temp:
        tmp=Path(temp)
        for page,group in enumerate(groups):
            if page:doc.add_page_break()
            if page==0:_header(doc,data,config,look,Path(case),tmp)
            elif look.get('repeat_name'):_para(doc,data['person']['name'],size=10,gray='777777',after=4)
            last_section=None
            for k,b in enumerate(group):
                source=first_section.get(b['id'])
                if source and (source!=last_section or k==0):
                    section_title,category=source
                    # Open sections without headings still get a proper heading when moved to a new page.
                    if section_title:_section(doc,section_title,category,look,tmp)
                    last_section=source
                _block(doc,b,look,tmp,False)
        if look.get('footer')=='page_number':
            p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            r=p.add_run();r.font.color.rgb=_rgb('888888');r.font.size=Pt(8)
            fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
        doc.core_properties.title='CV - '+data['person']['name'];doc.core_properties.author=data['person']['name'];doc.core_properties.comments=''
        doc.save(path)
